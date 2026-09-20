from pathlib import Path
import json
import xml.etree.ElementTree as ET
import openpyxl
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.graphics import renderSVG

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
for name in ['donnees', 'cartes', 'impression', 'templates', 'assets/armes', 'assets/factions', 'assets/barcodes', 'assets/races', 'assets/effets', 'assets/illustrations', 'assets/cristaux', 'verification']:
    (ROOT / name).mkdir(parents=True, exist_ok=True)

elements = [
    ('AERO','AIR','HYDRO','ELECTRO','89EDC2',155),
    ('HYDRO','EAU','PYRO','AERO','38CBED',192),
    ('ELECTRO','ELECTRICITE','AERO','MINERO','F5DC32',53),
    ('PYRO','FEU','CRYO','HYDRO','FF704A',12),
    ('CRYO','GLACE','HERBO','PYRO','ACDFFD',201),
    ('LUXO','LUMIERE','HEMATO','NECRO','F4E9AF',47),
    ('MINERO','ROCHE','ELECTRO','GEO','B3BFCA',210),
    ('HERBO','PLANTE','GEO','CRYO','89D355',100),
    ('HEMATO','SANG','NECRO','LUXO','E84976',343),
    ('NECRO','TENEBRES','LUXO','HEMATO','A979F7',270),
    ('GEO','TERRE','MINERO','HERBO','C7AB6D',41),
    ('RAINBOW','RAINBOW',None,None,'E4D5FB',280),
]
element_map = {x[0]:dict(zip(['id','label','strong_against','weak_against','color','hue'],x)) for x in elements}

def card(element, name, race, faction, job, title, positions, weapon, atk, defense, magic, barriers, text, source='proposition V1'):
    el = element_map[element]
    return dict(element=element,name=name,race=race,faction=faction,job=job,title=title,positions=positions,weapon=weapon,atk=atk,defense=defense,magic=magic,barriers=barriers,text=text,source=source,color=el['color'],hue=el['hue'],advantage=30,disadvantage=30)

cards = [
card('ELECTRO','MOMO','ROBOT','Chroma','NETTOYEUR','ROBOT ARTISTE MAGIQUE',[3,4],'Instrument',[202,167,'retry',84,'mana',29],[200,167,'retry',100,59,11],[6,5,3,1],[6,5,2],"2XM0 alias \u201cMomo\u201d est un petit robot triste vivant dans une décharge. Etant auparavant un petit robot de ménage, il est doté d'un talent artistique rare et incompris, il utilise une flûte traversière magique pour créer des sortilèges enchanteurs.",'carte JPG Momo et positions confirmees'),
card('AERO','CANA','FALCO','Nestown','MESSAGERE','LE DERNIER PASSAGE',[2,3],'Arc',[254,211,168,125,83,39],[203,170,136,101,66,31],[3],[4],"Le pont a cédé avant l'aube. Cana replie son aile blessée et tend le message au-delà du vide. A Nestown, une promesse doit parfois voyager plus loin que celui qui la porte."),
card('HYDRO','JULIENNE','SIRENA','Thalassea','GARDIENNE','UN SOUFFLE SOUS LA MAREE',[5],'Bâton',[177,149,118,88,'mana',28],[208,174,137,104,68,33],[6,5,4,3,1],[6,4],"Sous les portes noyées de Thalasséa, Julienne retient la mer au-dessus des naufragés. Chaque souffle qu'elle leur offre lui coûte un peu de force. Elle ne relâchera l'arche qu'au dernier survivant."),
card('PYRO','DARNAKO','DRAX','Vulkar','FORGERON','CE QUE LE FEU EPARGNE',[2],'Epée longue',[287,240,190,143,96,45],[176,144,113,86,55,27],[6,4],[5],"La forge s'effondre. Darnako cale sa lame sous une poutre brûlante et abrite le dernier apprenti. Aujourd'hui, sa force ne sert ni un roi ni une guerre : elle laisse à un enfant le temps de vivre."),
card('CRYO','MALINIA','HUMAIN','Niveria','VEILLEUSE','LA MEMOIRE SOUS LE GIVRE',[4],'Sceptre',[293,245,196,146,99,47],[178,147,116,87,57,26],[6,5,4,3,2,1],[6,3],"Au fond du lac dort un souvenir que Niveria croyait perdu. Malinia retient les fissures du gel pendant qu'une lueur ancienne réchauffe ses mains. Sauver le passé commence parfois par le laisser fondre."),
card('LUXO','AELIS','AURELION','Solaria','SOIGNEUSE','LA LUMIERE AUX OUBLIES',[5],'Bâton',[178,147,117,86,56,'revive'],[207,172,138,103,67,32],[6,5,4,3,2],[6,5,3],"Dans le sanctuaire brisé, Aelis ramasse les fragments d'une lanterne. Elle en recoud la lumière pour ceux que la cité a rejetés. Son premier miracle n'est pas de juger les âmes, mais de leur faire une place."),
card('MINERO','BALMHYR','NAIN','Durane','MENTOR','LA ROCHE QUI ELEVE',[1],'Hache',[208,186,165,121,99,77],[296,264,202,148,110,91],[6,4,2],[6,3],"Derrière sa maison de Canyonero, Balmhyr soulève la pierre sous les pas de ses compagnons. L'ancien exilé leur enseigne qu'un pouvoir ne vaut pas seulement par ce qu'il renverse, mais par ceux qu'il aide à se relever.",'statistiques Personnages ligne 13; scene inspiree du chapitre 8'),
card('HERBO','VICTORVINE','CARNIVERT','Arborium','VEILLEUR','LA DERNIERE POUSSE',[3],'Bâton',[231,195,154,115,75,36],[237,196,155,117,77,37],[6,4,2],[5,2],"Victorvine retire une jeune pousse des eaux mortes d'une usine. Derrière lui, les racines avancent déjà. Il connaît la patience de la forêt, et le prix qu'elle réclame lorsque les hommes l'ont trop longtemps ignorée."),
card('HEMATO','VERMINIA','VAMP','Draevenheim','DISSIDENTE','LE PRIX DU PASSAGE',[3],'Dague',[238,198,159,118,79,38],[233,194,'dodge',112,72,33],[6,5,3],[6,3],"Verminia ferme la grille derrière les serviteurs en fuite. Un fil de sang tient tête au fer de sa propre maison. Quand sa famille découvrira la trahison, elle aura déjà choisi à qui devait servir son héritage."),
card('NECRO','ZVIRI','SKULLZ','Cryptown','ASSASSIN','TUEUSE REDOUTABLE',[2],'Fouet',[291,240,'death',134,95,48],['dodge',140,107,84,57,29],[],[],"Mercenaire assassin du royaume des Skullz. Maîtresse des ombres, elle est engagée pour tuer et récupérer le sang de se ses victimes. Sa réputation de tueuse redoutable, implacable et sans pitié la précède, semant la terreur partout où elle passe.",'carte JPG Zviri; texte original preserve'),
card('GEO','NAZAR','RHINOZ','Zarok','GARDIEN','TENIR LA TERRE',[1],'Lance',[204,169,134,99,64,31],[296,247,197,146,97,47],[6,3],[4,2],"La pluie emporte la route de Zarok. Nazar enfonce ses mains dans la boue et relève la berge autour d'une charrette immobilisée. Il ne commande pas à la terre : il lui demande de tenir encore un instant."),
card('RAINBOW','KAYLIS','HUMAIN','Z13','COMBATTANTE','LES CENDRES ET LES COULEURS',[5],'Dague',[180,147,118,90,43,'revive'],[209,175,139,104,69,'retry'],[6,5,4,3,2],[6,3],"Dans le silence de la Tour de Fer, Kaylis pose ses bras tatoués sur le pupitre brisé. Les couleurs gagnent la pierre, mais rien n'efface l'absence de Lanio. Ce qui naît ici est une promesse, pas une victoire.",'ATK Personnages ligne 16; DEF proposition V1; scene inspiree des chapitres 16-17'),
]
for i,c in enumerate(cards,1):
    c['id'] = f'{i:08d}'
    c['slug'] = f"{i:02d}_{c['element']}_{c['name']}"
    c['art'] = str(ROOT/'assets'/'illustrations'/f"{c['element'].lower()}.png").replace('\\','/')
    c['reference'] = str(PROJECT/'main'/('V1_P1_FR_Momo (2).jpg' if c['name']=='MOMO' else 'FR_V1_Zviri.jpg')).replace('\\','/') if c['name'] in ['MOMO','ZVIRI'] else None
    if c['element']=='RAINBOW':
        c['advantage']=40;c['disadvantage']=0

def save(name,value):
    (ROOT/'donnees'/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf8')

save('cartes.json',cards)
save('elements.json',element_map)
save('registre_identifiants.json',[dict(id=c['id'],personnage=c['name'],version=c['title'],element=c['element'],fichier=c['slug']) for c in cards])
rules=dict(version='V1 - prototype',deck_size=10,positions={1:'Tank',2:'DPS physique',3:'Middle',4:'DPS magique',5:'Support / Healer'},rainbow_limit=1,weapon_modifier=50,barrier=30,synergy=[0,0,10,20,30,40],faction_stat='ATK',race_stat='DEF',ties='defender survives',special_effects='resolve separately; numeric model never treats an effect as zero')
save('regles.json',rules)

w = openpyxl.load_workbook(PROJECT/'main'/'Kalistar.xlsx')
weapons = w['Armes V2']
names = [weapons.cell(r,1).value for r in range(3,23)]
matrix = {}
for i,name in enumerate(names):
    matrix[name]={}
    for j,other in enumerate(names):
        c=weapons.cell(3+i,14+j)
        rgb=c.fill.fgColor.rgb if c.fill.fgColor.type=='rgb' else ''
        sign=-1 if rgb=='FFE6B8AF' else 1 if rgb=='FFB6D7A8' else 0
        matrix[name][other]=sign*int(c.value) if isinstance(c.value,(int,float)) else 0
for i,img in enumerate(weapons._images[:20]):
    (ROOT/'assets'/'armes'/f'{i:02d}.jpg').write_bytes(img._data())
for c in cards:
    c['weapon_index']=next(i for i,n in enumerate(names) if n.lower()==c['weapon'].lower())
save('cartes.json',cards)
save('armes.json',matrix)
for img in w['Ville']._images:
    name=w['Ville'].cell(2,img.anchor._from.col+1).value
    (ROOT/'assets'/'factions'/f'{name}.png').write_bytes(img._data())
for c in cards:
    drawing=createBarcodeDrawing('Code128',value=c['id'],barWidth=2,barHeight=22,humanReadable=False,quiet=True,lquiet=20,rquiet=20)
    raw=renderSVG.drawToString(drawing)
    tree=ET.fromstring(raw)
    body=''.join(ET.tostring(x,encoding='unicode') for x in tree)
    width=float(drawing.width)
    assert width<=210,(c['id'],width)
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" width="22" height="210" viewBox="0 0 22 210"><rect width="22" height="210" fill="white"/><g transform="translate(22 {(210-width)/2}) rotate(90)">{body}</g></svg>'
    (ROOT/'assets'/'barcodes'/f"{c['id']}.svg").write_text(svg,encoding='utf8')
print('Prepared',len(cards),'cards, signed weapon matrix, icons, flags and unique Code128 SVGs')
